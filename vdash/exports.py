"""Downloads: per-chart exports and full-list exports, as CSV or Excel.

* ``kpi_table`` turns any KPI result (time series, breakdown, heatmap, table)
  into rows + columns, so every chart can be downloaded as data.
* ``build_export_sql`` composes a full-list export from the catalog definition,
  keeping only the optional joins whose tables exist in the tenant warehouse.
* ``to_csv`` / ``to_xlsx`` write the file. Numbers stay numbers (no display
  formatting), CSV is UTF-8 with BOM so Excel shows Vietnamese/Chinese text
  correctly, and the Excel file carries an "About" sheet (tenant, filters, period).
"""
from __future__ import annotations

import csv
import io
import re
from datetime import datetime, timezone
from typing import Any

from .catalog import Export, Kpi
from .datasources import DataSource
from .engine import bucket_label
from .i18n import t

_ALIAS_RE = re.compile(r"\b([A-Za-z_][A-Za-z0-9_]*)\.")


# ---------------------------------------------------------------- KPI result -> table
def kpi_table(kpi: Kpi, res: dict, lang: str) -> tuple[list[dict], list[dict]]:
    """Return (columns, rows); columns = [{key, label, fmt}]. Labels/values already translated."""
    typ = res.get("type")
    if typ == "timeseries":
        cols = [{"key": "period", "label": t(lang, "Period"), "fmt": "text"}]
        cols += [{"key": f"s{i}", "label": s.get("name") or t(lang, kpi.title), "fmt": s.get("fmt") or kpi.fmt}
                 for i, s in enumerate(res["series"])]
        rows = []
        for j, x in enumerate(res["x"]):
            r = {"period": bucket_label(x, res.get("grain", "day")) if len(str(x)) == 10 else x}
            for i, s in enumerate(res["series"]):
                r[f"s{i}"] = s["data"][j]
            rows.append(r)
        return cols, rows
    if typ == "breakdown":
        cols = [{"key": "label", "label": t(lang, "Category"), "fmt": "text"},
                {"key": "value", "label": t(lang, kpi.title), "fmt": kpi.fmt}]
        return cols, [{"label": i["label"], "value": i["value"]} for i in res["items"]]
    if typ == "heatmap":
        cols = [{"key": "y", "label": "", "fmt": "text"}] + [{"key": f"c{i}", "label": x, "fmt": kpi.fmt}
                                                            for i, x in enumerate(res["x"])]
        rows = [{"y": y} for y in res["y"]]
        for xi, yi, v in res["cells"]:
            rows[yi][f"c{xi}"] = v
        return cols, rows
    if typ == "table":
        cols = [{"key": c.key, "label": t(lang, c.label), "fmt": c.fmt} for c in kpi.columns]
        return cols, res["rows"]
    if typ == "metric":
        cols = [{"key": "metric", "label": t(lang, "KPI"), "fmt": "text"},
                {"key": "value", "label": t(lang, "Value"), "fmt": kpi.fmt},
                {"key": "compare", "label": t(lang, "Comparison value"), "fmt": kpi.fmt}]
        rows = [{"metric": t(lang, kpi.title), "value": res.get("value"), "compare": res.get("compare_value")}]
        return cols, rows
    return [], []


# ---------------------------------------------------------------- full-list exports
def included_joins(exp: Export, ds: DataSource) -> list:
    cols = ds.columns()
    return [j for j in exp.joins if j.table in cols]


def active_filters(exp: Export, joins: list) -> list:
    aliases = {j.alias for j in joins}
    return [f for f in exp.filters if f.needs is None or f.needs in aliases]


def _from(exp: Export, joins: list) -> str:
    return exp.base_from + "".join(f" LEFT JOIN {j.table} {j.alias} ON {j.on}" for j in joins)


def build_export_sql(exp: Export, ds: DataSource, chosen: dict[str, str | None]) -> tuple[str, dict, list]:
    joins = included_joins(exp, ds)
    cols = exp.columns + [c for j in joins for c in j.columns]
    params: dict[str, Any] = {}
    where = [exp.base_where]
    for f in active_filters(exp, joins):
        val = chosen.get(f.param)
        params[f"f_{f.param}"] = val if val is not None else ""
        params[f"f_{f.param}_all"] = 0 if val else 1
        where.append(f"(:f_{f.param}_all = 1 OR {f.expr} = :f_{f.param})")
    aliases = {j.alias for j in joins} | {exp.base_from.split()[-1]}
    order = exp.order_by if set(_ALIAS_RE.findall(exp.order_by)) <= aliases else cols[0].expr
    sql = (f"SELECT {', '.join(f'{c.expr} AS {c.key}' for c in cols)} FROM {_from(exp, joins)} "
           f"WHERE {' AND '.join(where)} ORDER BY {order} LIMIT {int(exp.max_rows)}")
    return sql, params, cols


def build_count_sql(exp: Export, ds: DataSource, chosen: dict[str, str | None]) -> tuple[str, dict]:
    """v58: how many rows the list has for the chosen filters (the same conditions as ``build_export_sql``)."""
    joins = included_joins(exp, ds)
    params: dict[str, Any] = {}
    where = [exp.base_where]
    for f in active_filters(exp, joins):
        val = chosen.get(f.param)
        params[f"f_{f.param}"] = val if val is not None else ""
        params[f"f_{f.param}_all"] = 0 if val else 1
        where.append(f"(:f_{f.param}_all = 1 OR {f.expr} = :f_{f.param})")
    return f"SELECT COUNT(*) AS n FROM {_from(exp, joins)} WHERE {' AND '.join(where)}", params


def filter_options(exp: Export, ds: DataSource, tenant_id: str) -> list[dict]:
    joins = included_joins(exp, ds)
    out = []
    for f in active_filters(exp, joins):
        sql = (f"SELECT {f.expr} AS val, COUNT(*) AS n FROM {_from(exp, joins)} WHERE {exp.base_where} "
               f"GROUP BY {f.expr} ORDER BY n DESC")
        try:
            rows = [r for r in ds.query(sql, {"tenant_id": tenant_id}) if r["val"] not in (None, "")]
        except Exception:
            rows = []
        if len(rows) >= 2:  # a filter with one value is not a choice
            out.append({"param": f.param, "label": f.label,
                        "options": [{"value": str(r["val"]), "count": int(r["n"])} for r in rows]})
    return out


# ---------------------------------------------------------------- writers
_ROUND = {"currency": 2, "pct": 2, "ratio": 4, "number": 3, "days": 2, "hours": 2, "score": 2, "int": 0}


def _cell(v: Any, fmt: str) -> Any:
    """Round model noise (e.g. 9.2e-20) away; keep numbers numeric."""
    if isinstance(v, float) and fmt in _ROUND:
        v = round(v, _ROUND[fmt])
        return int(v) if fmt == "int" or v.is_integer() else v
    return v


def _label(c: dict) -> str:
    return f"{c['label']} (%)" if c.get("fmt") == "pct" and c.get("label") else c["label"]


def to_csv(cols: list[dict], rows: list[dict]) -> bytes:
    buf = io.StringIO()
    w = csv.writer(buf)
    w.writerow([_label(c) for c in cols])
    for r in rows:
        w.writerow(["" if r.get(c["key"]) is None else _cell(r.get(c["key"]), c["fmt"]) for c in cols])
    return ("﻿" + buf.getvalue()).encode("utf-8")


XLSX_FMT = {"currency": "#,##0", "number": "#,##0.##", "int": "#,##0", "pct": "0.0", "ratio": "0.00",
            "days": "0.0", "hours": "0.0", "score": "0.0"}


def to_xlsx(title: str, cols: list[dict], rows: list[dict], about: list[tuple[str, str]]) -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    ws = wb.active
    ws.title = re.sub(r"[\[\]:*?/\\]", " ", title)[:31] or "Data"
    head_fill = PatternFill("solid", fgColor="E6F0FC")
    for ci, c in enumerate(cols, 1):
        cell = ws.cell(row=1, column=ci, value=_label(c))
        cell.font = Font(bold=True)
        cell.fill = head_fill
        cell.alignment = Alignment(vertical="top", wrap_text=True)
    for ri, r in enumerate(rows, 2):
        for ci, c in enumerate(cols, 1):
            v = _cell(r.get(c["key"]), c["fmt"])
            cell = ws.cell(row=ri, column=ci, value=v)
            if c["fmt"] in XLSX_FMT and isinstance(v, (int, float)):
                cell.number_format = XLSX_FMT[c["fmt"]]
    for ci, c in enumerate(cols, 1):
        width = max([len(str(c["label"]))] + [len(str(r.get(c["key"]) or "")) for r in rows[:500]])
        ws.column_dimensions[get_column_letter(ci)].width = min(max(10, width + 2), 60)
    ws.freeze_panes = "A2"
    if rows:
        ws.auto_filter.ref = f"A1:{get_column_letter(len(cols))}{len(rows) + 1}"
    info = wb.create_sheet("About")
    for ri, (k, v) in enumerate(about + [("Generated (UTC)", datetime.now(timezone.utc).strftime("%Y-%m-%d %H:%M")),
                                         ("Rows", str(len(rows)))], 1):
        info.cell(row=ri, column=1, value=k).font = Font(bold=True)
        info.cell(row=ri, column=2, value=v)
    info.column_dimensions["A"].width = 24
    info.column_dimensions["B"].width = 80
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()


def safe_filename(*parts: str) -> str:
    s = "_".join(p for p in parts if p)
    s = re.sub(r"[^\w\-.]+", "-", s, flags=re.UNICODE).strip("-")
    return s[:120] or "export"
